import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


EXCEL_FILE = "010826.xlsm"
DATABASE_FILE = "salary_program.db"
EFFECTIVE_FROM = "2026-08"


CATEGORY_MAP = {
    "שכר טיפול": "טיפול",
    "שכר קבוצות": "קבוצה",
    "שכר ביטול": "ביטול",
    "שכר אבחון": "אבחון",
    "שכר מעקב": "מעקב",
    "שכר דוח": "דו״ח",
    'שכר ישב"צ': "ישיבת צוות",
    "שכר הדרכה": "הדרכה",
    "שעת הכנה": "הכנה",
    "שכר המלצה צהל": "המלצה לצה״ל",
}


FIXED_MONTHLY_SALARIES = {
    "סיגל מגן": 8000,
    "אריאלה יקוטי": 7500,
    "איריס אהרוני מצא": 6500,
    "ויקטור": 7000,
}


SKIP_SHEETS = {
    "ראשי",
    "Sheet1",
}


def clean_name(value):
    if value is None:
        return None

    value = str(value).strip()
    value = re.sub(r"^-+\s*מפגשים עבור\s*", "", value)
    value = re.sub(r"^-+", "", value)
    value = re.sub(r"-+$", "", value)
    value = value.strip()

    return value or None


def find_real_name(sheet):
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value

            if isinstance(value, str) and "מפגשים עבור" in value:
                name = clean_name(value)

                if name:
                    return name

    first_cell = sheet.cell(row=1, column=1).value

    if isinstance(first_cell, str) and first_cell.strip():
        return first_cell.strip()

    return None


def find_rates(sheet):
    rates = {}

    for row_number in range(1, min(sheet.max_row, 20) + 1):
        for column_number in range(1, min(sheet.max_column, 12) + 1):
            label = sheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if label not in CATEGORY_MAP:
                continue

            rate = sheet.cell(
                row=row_number + 1,
                column=column_number,
            ).value

            if isinstance(rate, (int, float)):
                rates[CATEGORY_MAP[label]] = float(rate)

    return rates


def create_tables(connection):
    cursor = connection.cursor()

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS therapist_rates (
            therapist_name TEXT NOT NULL,
            category TEXT NOT NULL,
            rate REAL NOT NULL,
            effective_from TEXT NOT NULL,
            PRIMARY KEY (
                therapist_name,
                category,
                effective_from
            )
        )
        """
    )

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS pay_structures (
            therapist_name TEXT NOT NULL,
            pay_structure TEXT NOT NULL,
            monthly_salary REAL NOT NULL DEFAULT 0,
            effective_from TEXT NOT NULL,
            PRIMARY KEY (
                therapist_name,
                effective_from
            )
        )
        """
    )

    connection.commit()


def save_rate(connection, therapist_name, category, rate):
    cursor = connection.cursor()

    cursor.execute(
        """
        INSERT INTO therapist_rates (
            therapist_name,
            category,
            rate,
            effective_from
        )
        VALUES (?, ?, ?, ?)
        ON CONFLICT(
            therapist_name,
            category,
            effective_from
        )
        DO UPDATE SET
            rate = excluded.rate
        """,
        (
            therapist_name,
            category,
            rate,
            EFFECTIVE_FROM,
        ),
    )


def save_pay_structure(
    connection,
    therapist_name,
    pay_structure,
    monthly_salary,
):
    cursor = connection.cursor()

    cursor.execute(
        """
        INSERT INTO pay_structures (
            therapist_name,
            pay_structure,
            monthly_salary,
            effective_from
        )
        VALUES (?, ?, ?, ?)
        ON CONFLICT(
            therapist_name,
            effective_from
        )
        DO UPDATE SET
            pay_structure = excluded.pay_structure,
            monthly_salary = excluded.monthly_salary
        """,
        (
            therapist_name,
            pay_structure,
            monthly_salary,
            EFFECTIVE_FROM,
        ),
    )


excel_path = Path(EXCEL_FILE)
database_path = Path(DATABASE_FILE)

if not excel_path.exists():
    print()
    print("ERROR:")
    print(f"Could not find {EXCEL_FILE}")
    raise SystemExit

if not database_path.exists():
    print()
    print("ERROR:")
    print(f"Could not find {DATABASE_FILE}")
    raise SystemExit


workbook = load_workbook(
    EXCEL_FILE,
    data_only=True,
    keep_vba=True,
)

connection = sqlite3.connect(DATABASE_FILE)

create_tables(connection)

imported = []
needs_review = []


for sheet_name in workbook.sheetnames:

    if sheet_name in SKIP_SHEETS:
        continue

    sheet = workbook[sheet_name]
    therapist_name = find_real_name(sheet)

    if not therapist_name:
        needs_review.append(
            (sheet_name, "Could not identify therapist name")
        )
        continue

    rates = find_rates(sheet)

    if therapist_name in FIXED_MONTHLY_SALARIES:
        monthly_salary = FIXED_MONTHLY_SALARIES[therapist_name]

        save_pay_structure(
            connection,
            therapist_name,
            "Fixed monthly salary",
            monthly_salary,
        )

        imported.append(
            (therapist_name, "FIXED", monthly_salary)
        )

    else:
        save_pay_structure(
            connection,
            therapist_name,
            "Per activity",
            0,
        )

        saved_any_rate = False

        for category, rate in rates.items():
            save_rate(
                connection,
                therapist_name,
                category,
                rate,
            )
            saved_any_rate = True

        if saved_any_rate:
            imported.append(
                (therapist_name, "PER ACTIVITY", rates)
            )
        else:
            needs_review.append(
                (therapist_name, "No usable pay rates found")
            )


connection.commit()
connection.close()


print()
print("=" * 70)
print("AUGUST SALARY IMPORT FINISHED")
print("=" * 70)

print()
print("IMPORTED:")
print()

for name, pay_type, details in imported:
    print(name)
    print(f"  Pay structure: {pay_type}")

    if pay_type == "FIXED":
        print(f"  Monthly salary: {details} NIS")
    else:
        for category, rate in details.items():
            print(f"  {category}: {rate} NIS")

    print()


print("=" * 70)
print("NEEDS REVIEW:")
print("=" * 70)
print()

if needs_review:
    for name, reason in needs_review:
        print(f"{name}: {reason}")
else:
    print("None")

print()
print("=" * 70)
print("Employment status was NOT changed.")
print("Existing Employee/Freelance settings were preserved.")
print("=" * 70)